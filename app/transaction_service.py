"""Transaction monitoring: marketplace inquiries, staff follow-up, notifications, Excel import/export."""

from __future__ import annotations

import re
from datetime import date, datetime, time
from io import BytesIO

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import String, and_, cast, func, or_
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import joinedload

from app import db
from app.config import (
    ALL_LEAD_STATUSES,
    MARKETPLACE_CATEGORY_PRODUCTS,
    MARKETPLACE_LEAD_CLOSED_STATUSES,
    MARKETPLACE_LEAD_STATUS_COMPLETED,
    MARKETPLACE_LEAD_STATUS_FOR_CONTRACTOR,
    MARKETPLACE_LEAD_STATUS_FOR_FOLLOW_UP,
    MARKETPLACE_LEAD_STATUS_FOR_QUOTATION,
    MARKETPLACE_LEAD_STATUS_FOR_RESEARCH,
    MARKETPLACE_LEAD_STATUS_LABELS,
    MARKETPLACE_LEAD_STATUS_NEW,
    MARKETPLACE_LEAD_STATUS_ON_HOLD,
    MARKETPLACE_LEAD_STATUS_QUOTATION_IN_PROCESS,
    MARKETPLACE_LEAD_STATUS_QUOTATION_SUBMITTED,
    MARKETPLACE_LEAD_STATUS_TERMINATED,
    MARKETPLACE_LEAD_STATUSES,
    DUTY_PROJECT_COORDINATOR,
    PRODUCT_LEAD_STATUSES,
    PROJECT_LEAD_STATUSES,
    TRANSACTION_NOTE_KINDS,
    TRANSACTION_REF_PREFIXES,
    TRANSACTION_SOURCE_IMPORT,
    TRANSACTION_SOURCE_LABELS,
    TRANSACTION_SOURCE_LEGACY,
    TRANSACTION_SOURCE_MANUAL,
    TRANSACTION_SOURCE_WEB,
    TRANSACTION_TYPE_LABELS,
    TRANSACTION_TYPE_PROJECTS,
    USER_ROLE_ADMIN,
    USER_ROLE_PORTAL_ADMIN,
    USER_ROLE_STAFF,
    is_staff_or_admin,
    member_duties_for_type,
    team_slot_labels_for_type,
    team_slots_for_type,
)
from app.duty_service import (
    assignee_choices,
    can_be_assigned,
    duty_holders,
    duty_label,
    match_coordinator,
    staff_users,
    team_filter,
)
from app.models import (
    Contractor,
    MarketplaceLead,
    MarketplaceLeadHistory,
    MarketplaceLeadRead,
    MarketplaceListing,
    Member,
    Supplier,
    User,
)
from app.timeutil import manila_now, manila_today

EMAIL_RE = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")
ASSIGNABLE_ROLES = (USER_ROLE_PORTAL_ADMIN, USER_ROLE_ADMIN, USER_ROLE_STAFF)

# Editable free-text fields: name -> (label, max length or None for TEXT).
TRANSACTION_TEXT_FIELDS = {
    "item_name": ("Product / transaction", 255),
    "estimated_implementation": ("Estimated date of implementation", 160),
    "quantity": ("Qty / volume", None),
    "specifications": ("Specifications", None),
    "delivery_location": ("Delivered to / location", None),
    "referrer_name": ("ProF (referred by)", 255),
    "referrer_phone": ("ProF contact number", 120),
    "referrer_email": ("ProF email", 255),
    "client_company": ("Client / company", None),
    "guest_name": ("Client representative", 160),
    "guest_phone": ("Client contact number", 160),
    "guest_email": ("Client email", 255),
    "status_detail": ("Status details", None),
    "action_needed": ("Action needed", None),
    "project_coordinator": ("Project coordinator", 160),
    "assigned_contractor": ("Assigned contractor / supplier", 255),
    "remarks": ("Remarks", None),
    "message": ("Inquiry message", None),
}
TRANSACTION_DATE_FIELDS = {
    "date_requested": "Date requested",
    "date_completed": "Date completed",
    "delivery_date": "Scheduled delivery",
}
PROJECT_FIELD_LABELS = {
    "item_name": "Project / transaction",
    "specifications": "Project details and specifications",
    "delivery_location": "Location",
    "assigned_contractor": "Assigned contractor(s)",
}


def transaction_field_labels(transaction_type=None):
    labels = {field: label for field, (label, _max_len) in TRANSACTION_TEXT_FIELDS.items()}
    labels.update(TRANSACTION_DATE_FIELDS)
    if transaction_type == TRANSACTION_TYPE_PROJECTS:
        labels.update(PROJECT_FIELD_LABELS)
    return labels


# --------------------------------------------------------------------------- labels / numbering


def transaction_status_label(status):
    return MARKETPLACE_LEAD_STATUS_LABELS.get(status or MARKETPLACE_LEAD_STATUS_NEW, status or "New")


def transaction_status_choices(transaction_type=None):
    """Workflow statuses offered for a transaction type (projects add delivery stages, products order stages)."""
    if transaction_type == TRANSACTION_TYPE_PROJECTS:
        return PROJECT_LEAD_STATUSES
    if transaction_type == MARKETPLACE_CATEGORY_PRODUCTS:
        return PRODUCT_LEAD_STATUSES
    if transaction_type == "all":
        return ALL_LEAD_STATUSES
    return MARKETPLACE_LEAD_STATUSES


def set_transaction_status(lead, status, *, user_id=None, when=None):
    """Apply a status change and its side effects (time-in-status, project stage plan)."""
    now = when or manila_now()
    lead.status = status
    lead.status_changed_at = now
    if status in MARKETPLACE_LEAD_CLOSED_STATUSES and not lead.date_completed:
        lead.date_completed = now.date()
    if lead.transaction_type == TRANSACTION_TYPE_PROJECTS:
        from app.project_service import on_project_status_change

        on_project_status_change(lead, status, now.date(), user_id=user_id)
    elif lead.transaction_type == MARKETPLACE_CATEGORY_PRODUCTS:
        from app.product_service import on_product_status_change

        on_product_status_change(lead, status, user_id=user_id)


def transaction_type_label(transaction_type):
    key = transaction_type or MARKETPLACE_CATEGORY_PRODUCTS
    return TRANSACTION_TYPE_LABELS.get(key, key.replace("_", " ").title())


def transaction_source_label(source):
    return TRANSACTION_SOURCE_LABELS.get(source or TRANSACTION_SOURCE_WEB, source or "")


def format_reference_number(transaction_type, inquiry_no):
    prefix = TRANSACTION_REF_PREFIXES.get(transaction_type or MARKETPLACE_CATEGORY_PRODUCTS, "TRX")
    return f"{prefix}-{int(inquiry_no):04d}"


def next_inquiry_no(transaction_type):
    current = (
        db.session.query(func.max(MarketplaceLead.inquiry_no))
        .filter(MarketplaceLead.transaction_type == transaction_type)
        .scalar()
    )
    return int(current or 0) + 1


def preview_reference_number(transaction_type):
    """Next reference number (final number is allocated on submit)."""
    return format_reference_number(transaction_type, next_inquiry_no(transaction_type))


def _clean(value, max_len=None):
    if value is None:
        return None
    text = str(value).strip()
    if not text:
        return None
    if max_len:
        text = text[:max_len]
    return text


def _parse_form_date(raw, label):
    text = (raw or "").strip()
    if not text:
        return None
    try:
        return datetime.strptime(text[:10], "%Y-%m-%d").date()
    except ValueError as exc:
        raise ValueError(f"{label} must be a valid date (YYYY-MM-DD).") from exc


def _append_history(lead, *, event_type, note=None, user_id=None):
    entry = MarketplaceLeadHistory(
        lead_id=lead.lead_id,
        event_type=event_type,
        status=lead.status,
        note=(note or "").strip() or None,
        created_by_user_id=user_id,
        created_at=manila_now(),
    )
    db.session.add(entry)
    return entry


def _insert_with_reference(lead):
    """Allocate the next inquiry number for the lead's type, retrying on concurrent inserts."""
    for _ in range(5):
        lead.inquiry_no = next_inquiry_no(lead.transaction_type)
        lead.reference_number = format_reference_number(lead.transaction_type, lead.inquiry_no)
        try:
            with db.session.begin_nested():
                db.session.add(lead)
                db.session.flush()
            return lead
        except IntegrityError:
            continue
    raise ValueError("Could not allocate a transaction reference number. Please try again.")


# --------------------------------------------------------------------------- creation


def _member_contact(member):
    if member is None:
        return None, None, None
    return member.full_name, getattr(member, "phone", None), getattr(member, "email", None)


def create_inquiry(
    *,
    listing=None,
    transaction_type=None,
    item_name=None,
    guest_name=None,
    guest_phone=None,
    guest_email=None,
    client_company=None,
    quantity=None,
    specifications=None,
    delivery_location=None,
    estimated_implementation=None,
    message=None,
    source_path=None,
    attributed_member_id=None,
):
    """Public inquiry from a guest or member. Full name, email and contact number are required."""
    name = _clean(guest_name, 160)
    email = _clean(guest_email, 255)
    phone = _clean(guest_phone, 160)
    if not name:
        raise ValueError("Full name is required.")
    if not email:
        raise ValueError("Email address is required.")
    if not EMAIL_RE.match(email):
        raise ValueError("Enter a valid email address.")
    if not phone:
        raise ValueError("Contact number is required.")
    if len(re.sub(r"\D", "", phone)) < 7:
        raise ValueError("Enter a valid contact number.")

    category = transaction_type or (listing.category if listing is not None else None)
    category = category or MARKETPLACE_CATEGORY_PRODUCTS
    member = db.session.get(Member, attributed_member_id) if attributed_member_id else None
    ref_name, ref_phone, ref_email = _member_contact(member)
    now = manila_now()
    title = listing.title if listing is not None else item_name
    if not _clean(title) and category == TRANSACTION_TYPE_PROJECTS:
        title = "Project inquiry"

    lead = MarketplaceLead(
        listing_id=listing.listing_id if listing is not None else None,
        interest_category=category,
        transaction_type=category,
        item_name=_clean(title, 255),
        attributed_member_id=member.member_id if member else None,
        referrer_name=_clean(ref_name, 255),
        referrer_phone=_clean(ref_phone, 120),
        referrer_email=_clean(ref_email, 255),
        guest_name=name,
        guest_phone=phone,
        guest_email=email,
        client_company=_clean(client_company),
        quantity=_clean(quantity),
        specifications=_clean(specifications),
        delivery_location=_clean(delivery_location),
        estimated_implementation=_clean(estimated_implementation, 160),
        message=_clean(message),
        source_path=_clean(source_path, 255),
        status=MARKETPLACE_LEAD_STATUS_NEW,
        source=TRANSACTION_SOURCE_WEB,
        date_requested=now.date(),
        created_at=now,
        updated_at=now,
        status_changed_at=now,
    )
    _insert_with_reference(lead)
    _append_history(
        lead,
        event_type="created",
        note=f"{transaction_type_label(category)} inquiry received via website",
    )
    db.session.commit()
    return lead


def _apply_text_fields(lead, form, changes):
    labels = transaction_field_labels(lead.transaction_type)
    for field, (_label, max_len) in TRANSACTION_TEXT_FIELDS.items():
        if field not in form:
            continue
        value = _clean(form.get(field), max_len)
        if value != (getattr(lead, field) or None):
            setattr(lead, field, value)
            changes.append(f"{labels[field]} updated")


def _apply_date_fields(lead, form, changes):
    for field, label in TRANSACTION_DATE_FIELDS.items():
        if field not in form:
            continue
        value = _parse_form_date(form.get(field), label)
        if value != getattr(lead, field):
            setattr(lead, field, value)
            changes.append(f"{label} -> {value.isoformat() if value else 'cleared'}")


def _resolve_assignee(raw, transaction_type=None, duty=None):
    text = (raw or "").strip()
    if not text:
        return None
    if not text.isdigit():
        raise ValueError("Invalid staff assignment.")
    user = db.session.get(User, int(text))
    if not can_be_assigned(user, transaction_type, duty):
        if duty and duty in member_duties_for_type(transaction_type):
            raise ValueError(
                f"Choose an Admin or Staff user, or a member holding the {duty_label(duty)} duty."
            )
        raise ValueError("Transactions can only be assigned to Admin or Staff users.")
    return user


def _display_name(user):
    return (user.full_name or user.username) if user else "nobody"


def _apply_team(lead, form, changes, acting_user_id):
    """Team slots for the transaction type: project team, product team, or just the assigned staff."""
    slots = team_slots_for_type(lead.transaction_type)
    labels = team_slot_labels_for_type(lead.transaction_type)
    has_team = len(slots) > 1
    for slot, duty in slots.items():
        if slot not in form:
            continue
        user = _resolve_assignee(form.get(slot), lead.transaction_type, duty)
        new_id = user.user_id if user else None
        if new_id == getattr(lead, slot):
            continue
        setattr(lead, slot, new_id)
        changes.append(f"{labels[slot]}: {_display_name(user)}" if has_team else f"Assigned to {_display_name(user)}")
        if lead.transaction_type == TRANSACTION_TYPE_PROJECTS and slot == "assigned_user_id" and user is not None:
            lead.project_coordinator = _display_name(user)[:160]
        if lead.lead_id:
            _notify_assignee(lead, new_id, acting_user_id)


def _apply_contractor(lead, form, changes):
    """Staff assign a registered contractor to a project; it fills the contractor name and opens the contractor portal."""
    if lead.transaction_type != TRANSACTION_TYPE_PROJECTS or "awarded_contractor_id" not in form:
        return
    text = (form.get("awarded_contractor_id") or "").strip()
    contractor = None
    if text:
        if not text.isdigit():
            raise ValueError("Choose the contractor from the list.")
        contractor = db.session.get(Contractor, int(text))
        if contractor is None:
            raise ValueError("Contractor not found.")
    new_id = contractor.contractor_id if contractor else None
    if new_id == lead.awarded_contractor_id:
        return
    previous = db.session.get(Contractor, lead.awarded_contractor_id) if lead.awarded_contractor_id else None
    lead.awarded_contractor_id = new_id
    if contractor is not None:
        if contractor.company_name.lower() not in (lead.assigned_contractor or "").lower():
            lead.assigned_contractor = contractor.company_name[:255]
        changes.append(f"Contractor: {contractor.company_name}")
    else:
        if previous is not None and lead.assigned_contractor == previous.company_name:
            lead.assigned_contractor = None
        changes.append(f"Contractor removed{f' ({previous.company_name})' if previous else ''}")


def _apply_supplier(lead, form, changes):
    """Staff assign a registered supplier to a product order; it fills the supplier name and opens the supplier portal."""
    if lead.transaction_type != MARKETPLACE_CATEGORY_PRODUCTS or "supplier_id" not in form:
        return
    text = (form.get("supplier_id") or "").strip()
    supplier = None
    if text:
        if not text.isdigit():
            raise ValueError("Choose the supplier from the list.")
        supplier = db.session.get(Supplier, int(text))
        if supplier is None:
            raise ValueError("Supplier not found.")
    new_id = supplier.supplier_id if supplier else None
    if new_id == lead.supplier_id:
        return
    previous = db.session.get(Supplier, lead.supplier_id) if lead.supplier_id else None
    lead.supplier_id = new_id
    if supplier is not None:
        if supplier.company_name.lower() not in (lead.assigned_contractor or "").lower():
            lead.assigned_contractor = supplier.company_name[:255]
        changes.append(f"Supplier: {supplier.company_name}")
    else:
        if previous is not None and lead.assigned_contractor == previous.company_name:
            lead.assigned_contractor = None
        changes.append(f"Supplier removed{f' ({previous.company_name})' if previous else ''}")


def supplier_options():
    return [
        {"id": s.supplier_id, "name": s.company_name}
        for s in Supplier.query.order_by(Supplier.company_name.asc()).all()
    ]


def _resolve_member(raw):
    text = (raw or "").strip()
    if not text:
        return None
    if not text.isdigit():
        raise ValueError("Referred-by member ID must be a number.")
    member = db.session.get(Member, int(text))
    if not member:
        raise ValueError(f"Member #{text} was not found.")
    return member


def _notify_assignee(lead, assignee_id, acting_user_id):
    """Make the transaction unread for a newly assigned user (unless they assigned themselves)."""
    if not assignee_id or assignee_id == acting_user_id:
        return
    MarketplaceLeadRead.query.filter_by(lead_id=lead.lead_id, user_id=assignee_id).delete()


def create_manual_transaction(form, user):
    """Staff/Admin logs an offline inquiry (phone, walk-in, chat)."""
    transaction_type = (form.get("transaction_type") or MARKETPLACE_CATEGORY_PRODUCTS).strip()
    if transaction_type not in TRANSACTION_TYPE_LABELS:
        raise ValueError("Invalid transaction type.")
    listing = None
    listing_raw = (form.get("listing_id") or "").strip()
    if listing_raw.isdigit():
        listing = db.session.get(MarketplaceListing, int(listing_raw))

    now = manila_now()
    lead = MarketplaceLead(
        transaction_type=transaction_type,
        interest_category=transaction_type,
        listing_id=listing.listing_id if listing else None,
        status=MARKETPLACE_LEAD_STATUS_NEW,
        source=TRANSACTION_SOURCE_MANUAL,
        created_at=now,
        updated_at=now,
    )
    _apply_text_fields(lead, form, [])
    _apply_date_fields(lead, form, [])
    if not lead.item_name and listing:
        lead.item_name = listing.title
    if not lead.item_name:
        raise ValueError(f"{transaction_field_labels(transaction_type)['item_name']} is required.")
    if not lead.guest_name and not lead.client_company:
        raise ValueError("Client representative or company is required.")
    if not lead.guest_phone and not lead.guest_email:
        raise ValueError("Provide the client's contact number or email.")
    lead.date_requested = lead.date_requested or now.date()

    status = (form.get("status") or MARKETPLACE_LEAD_STATUS_NEW).strip()
    if status not in transaction_status_choices(transaction_type):
        raise ValueError("Invalid status.")
    lead.status = status
    lead.status_changed_at = now
    member = _resolve_member(form.get("attributed_member_id"))
    if member:
        lead.attributed_member_id = member.member_id
        if not lead.referrer_name:
            lead.referrer_name, phone, email = _member_contact(member)
            lead.referrer_phone = lead.referrer_phone or _clean(phone, 120)
            lead.referrer_email = lead.referrer_email or _clean(email, 255)
    _apply_team(lead, form, [], user.user_id)
    _apply_contractor(lead, form, [])
    _apply_supplier(lead, form, [])
    if lead.is_closed and not lead.date_completed:
        lead.date_completed = now.date()

    _insert_with_reference(lead)
    if transaction_type == TRANSACTION_TYPE_PROJECTS:
        from app.project_service import on_project_status_change

        on_project_status_change(lead, status, now.date(), user_id=user.user_id)
    _append_history(lead, event_type="created", note="Transaction logged manually", user_id=user.user_id)
    mark_transaction_read(lead.lead_id, user.user_id, commit=False)
    db.session.commit()
    return lead


# --------------------------------------------------------------------------- staff updates


def update_transaction(lead, form, user):
    """Apply the staff edit form; returns True when something changed (history row written)."""
    changes = []
    _apply_text_fields(lead, form, changes)
    _apply_date_fields(lead, form, changes)

    if "status" in form:
        status = (form.get("status") or "").strip()
        if status not in transaction_status_choices(lead.transaction_type):
            raise ValueError("Invalid status.")
        if status != (lead.status or MARKETPLACE_LEAD_STATUS_NEW):
            changes.insert(0, f"Status: {transaction_status_label(lead.status)} -> {transaction_status_label(status)}")
            set_transaction_status(lead, status, user_id=user.user_id)

    _apply_team(lead, form, changes, user.user_id)
    _apply_contractor(lead, form, changes)
    _apply_supplier(lead, form, changes)

    if "attributed_member_id" in form:
        member = _resolve_member(form.get("attributed_member_id"))
        new_id = member.member_id if member else None
        if new_id != lead.attributed_member_id:
            lead.attributed_member_id = new_id
            changes.append(f"Referred-by member -> {'#' + str(new_id) if new_id else 'none'}")

    note = (form.get("note") or "").strip()
    if not changes and not note:
        return False

    lead.updated_at = manila_now()
    summary = "; ".join(changes)
    if summary and note:
        summary = f"{summary}. {note}"
    _append_history(lead, event_type="update", note=summary or note, user_id=user.user_id)
    db.session.commit()
    return True


def add_followup_note(lead, kind, note, user):
    kind = (kind or "note").strip().lower()
    if kind not in TRANSACTION_NOTE_KINDS:
        raise ValueError("Invalid follow-up type.")
    text = (note or "").strip()
    if not text:
        if kind == "note":
            raise ValueError("Enter a note.")
        text = "Called the client" if kind == "call" else "Emailed the client"
    lead.updated_at = manila_now()
    _append_history(lead, event_type=kind, note=text, user_id=user.user_id)
    db.session.commit()


def assign_to_user(lead, user):
    if user.role not in ASSIGNABLE_ROLES:
        raise ValueError("Only Admin or Staff users can take transactions.")
    if lead.assigned_user_id == user.user_id:
        return False
    lead.assigned_user_id = user.user_id
    if lead.transaction_type == TRANSACTION_TYPE_PROJECTS:
        lead.project_coordinator = _display_name(user)[:160]
    lead.updated_at = manila_now()
    _append_history(
        lead,
        event_type="update",
        note=f"Assigned to {user.full_name or user.username} (self-assigned)",
        user_id=user.user_id,
    )
    db.session.commit()
    return True


def get_transaction(lead_id):
    return (
        MarketplaceLead.query
        .options(
            joinedload(MarketplaceLead.listing),
            joinedload(MarketplaceLead.attributed_member),
            joinedload(MarketplaceLead.assigned_user),
        )
        .filter(MarketplaceLead.lead_id == int(lead_id))
        .first()
    )


def transaction_history(lead_id):
    return (
        MarketplaceLeadHistory.query
        .options(joinedload(MarketplaceLeadHistory.created_by))
        .filter_by(lead_id=int(lead_id))
        .order_by(MarketplaceLeadHistory.created_at.desc(), MarketplaceLeadHistory.history_id.desc())
        .all()
    )


def assignee_options():
    return staff_users()


def team_options(transaction_type, lead=None):
    """Assignee choices per team slot (duty holders, or Staff/Admin until duties are set)."""
    return {
        slot: assignee_choices(duty, include_ids=[getattr(lead, slot)] if lead is not None else [])
        for slot, duty in team_slots_for_type(transaction_type).items()
        if duty
    }


# --------------------------------------------------------------------------- notifications


def _unread_filter(user_id, include_unassigned=True):
    already_read = (
        db.session.query(MarketplaceLeadRead.read_id)
        .filter(
            MarketplaceLeadRead.lead_id == MarketplaceLead.lead_id,
            MarketplaceLeadRead.user_id == user_id,
        )
        .exists()
    )
    scope = team_filter(user_id)
    if include_unassigned:
        scope = or_(
            and_(
                MarketplaceLead.source == TRANSACTION_SOURCE_WEB,
                MarketplaceLead.assigned_user_id.is_(None),
            ),
            scope,
        )
    return and_(
        ~already_read,
        MarketplaceLead.status.notin_(MARKETPLACE_LEAD_CLOSED_STATUSES),
        scope,
    )


def field_unread_count(user_id):
    """Unread projects for a member field agent (team slots only)."""
    if not user_id:
        return 0
    return MarketplaceLead.query.filter(_unread_filter(user_id, include_unassigned=False)).count()


def field_unread_ids(user_id):
    return {
        row.lead_id
        for row in MarketplaceLead.query.filter(_unread_filter(user_id, include_unassigned=False))
        .with_entities(MarketplaceLead.lead_id)
    }


def unread_transaction_count(user_id, role, transaction_type=None, exclude_type=None):
    if not user_id or not is_staff_or_admin(role):
        return 0
    query = MarketplaceLead.query.filter(_unread_filter(user_id))
    if transaction_type:
        query = query.filter(MarketplaceLead.transaction_type == transaction_type)
    if exclude_type:
        query = query.filter(MarketplaceLead.transaction_type != exclude_type)
    return query.count()


def unread_transaction_ids(user_id):
    return {
        row.lead_id
        for row in MarketplaceLead.query.filter(_unread_filter(user_id)).with_entities(MarketplaceLead.lead_id)
    }


def dashboard_transaction_alerts(user_id, role, limit=5):
    if not is_staff_or_admin(role):
        return None
    query = (
        MarketplaceLead.query
        .options(joinedload(MarketplaceLead.attributed_member))
        .filter(_unread_filter(user_id))
    )
    return {
        "unread_count": query.count(),
        "items": query.order_by(MarketplaceLead.created_at.desc()).limit(limit).all(),
    }


def mark_transaction_read(lead_id, user_id, commit=True):
    exists = MarketplaceLeadRead.query.filter_by(lead_id=lead_id, user_id=user_id).first()
    if exists:
        return
    db.session.add(MarketplaceLeadRead(lead_id=lead_id, user_id=user_id, read_at=manila_now()))
    if commit:
        try:
            db.session.commit()
        except IntegrityError:
            db.session.rollback()


def mark_all_transactions_read(user_id):
    ids = unread_transaction_ids(user_id)
    now = manila_now()
    for lead_id in ids:
        db.session.add(MarketplaceLeadRead(lead_id=lead_id, user_id=user_id, read_at=now))
    db.session.commit()
    return len(ids)


# --------------------------------------------------------------------------- monitoring list


def transaction_filters_from_args(args):
    def _text(key):
        return (args.get(key) or "").strip()

    return {
        "q": _text("q"),
        "transaction_type": _text("type") or MARKETPLACE_CATEGORY_PRODUCTS,
        "status": _text("status"),
        "stage": _text("stage"),
        "assigned": _text("assigned"),
        "source": _text("source"),
        "date_from": _text("date_from"),
        "date_to": _text("date_to"),
        "attention": _text("attention"),
    }


def search_transactions(filters, user=None, limit=1000):
    query = (
        MarketplaceLead.query
        .options(
            joinedload(MarketplaceLead.listing),
            joinedload(MarketplaceLead.attributed_member),
            joinedload(MarketplaceLead.assigned_user),
            joinedload(MarketplaceLead.estimator_user),
            joinedload(MarketplaceLead.site_engineer_user),
            joinedload(MarketplaceLead.sourcing_user),
            joinedload(MarketplaceLead.logistics_user),
            joinedload(MarketplaceLead.agent_user),
            joinedload(MarketplaceLead.supplier),
        )
    )
    transaction_type = filters.get("transaction_type")
    if transaction_type and transaction_type != "all":
        query = query.filter(MarketplaceLead.transaction_type == transaction_type)

    status = filters.get("status")
    if status in ALL_LEAD_STATUSES:
        query = query.filter(MarketplaceLead.status == status)
    stage = filters.get("stage")
    if stage == "open":
        query = query.filter(MarketplaceLead.status.notin_(MARKETPLACE_LEAD_CLOSED_STATUSES))
    elif stage == "closed":
        query = query.filter(MarketplaceLead.status.in_(MARKETPLACE_LEAD_CLOSED_STATUSES))

    assigned = filters.get("assigned")
    if assigned == "me" and user is not None:
        query = query.filter(team_filter(user.user_id))
    elif assigned == "unassigned":
        query = query.filter(MarketplaceLead.assigned_user_id.is_(None))
    elif assigned and assigned.isdigit():
        query = query.filter(team_filter(int(assigned)))

    source = filters.get("source")
    if source in TRANSACTION_SOURCE_LABELS:
        query = query.filter(MarketplaceLead.source == source)

    for key, op in (("date_from", "ge"), ("date_to", "le")):
        raw = filters.get(key)
        if not raw:
            continue
        try:
            value = datetime.strptime(raw[:10], "%Y-%m-%d").date()
        except ValueError:
            continue
        column = func.coalesce(MarketplaceLead.date_requested, func.date(MarketplaceLead.created_at))
        query = query.filter(column >= value if op == "ge" else column <= value)

    needle = filters.get("q")
    if needle:
        like = f"%{needle}%"
        query = query.outerjoin(
            MarketplaceListing, MarketplaceLead.listing_id == MarketplaceListing.listing_id
        ).filter(or_(
            MarketplaceLead.reference_number.ilike(like),
            cast(MarketplaceLead.inquiry_no, String).ilike(like),
            MarketplaceLead.item_name.ilike(like),
            MarketplaceListing.title.ilike(like),
            MarketplaceLead.guest_name.ilike(like),
            MarketplaceLead.guest_phone.ilike(like),
            MarketplaceLead.guest_email.ilike(like),
            MarketplaceLead.client_company.ilike(like),
            MarketplaceLead.referrer_name.ilike(like),
            MarketplaceLead.specifications.ilike(like),
            MarketplaceLead.delivery_location.ilike(like),
            MarketplaceLead.status_detail.ilike(like),
            MarketplaceLead.project_coordinator.ilike(like),
            MarketplaceLead.assigned_contractor.ilike(like),
        ))

    return (
        query
        .order_by(
            MarketplaceLead.inquiry_no.desc().nullslast(),
            MarketplaceLead.created_at.desc(),
        )
        .limit(limit)
        .all()
    )


def status_counts(transaction_type=None):
    query = db.session.query(MarketplaceLead.status, func.count(MarketplaceLead.lead_id))
    if transaction_type and transaction_type != "all":
        query = query.filter(MarketplaceLead.transaction_type == transaction_type)
    counts = {status: 0 for status in transaction_status_choices(transaction_type)}
    for status, count in query.group_by(MarketplaceLead.status).all():
        counts[status or MARKETPLACE_LEAD_STATUS_NEW] = counts.get(status, 0) + int(count)
    return counts


# --------------------------------------------------------------------------- Excel import


# Row fill colors used on the TBGP monitoring sheet (legend: RED = completed, BLACK = dead).
SHEET_COLOR_STATUS = {
    "FFFF0000": MARKETPLACE_LEAD_STATUS_COMPLETED,
    "theme1": MARKETPLACE_LEAD_STATUS_TERMINATED,
    "FF000000": MARKETPLACE_LEAD_STATUS_TERMINATED,
    "FFFFFF00": MARKETPLACE_LEAD_STATUS_FOR_FOLLOW_UP,
    "FF7030A0": MARKETPLACE_LEAD_STATUS_FOR_FOLLOW_UP,
    "FF26DC18": MARKETPLACE_LEAD_STATUS_FOR_QUOTATION,
    "FF2DFF8C": MARKETPLACE_LEAD_STATUS_FOR_QUOTATION,
    "FF00B0F0": MARKETPLACE_LEAD_STATUS_ON_HOLD,
    "FFFFC000": MARKETPLACE_LEAD_STATUS_FOR_RESEARCH,
    "theme5": MARKETPLACE_LEAD_STATUS_FOR_RESEARCH,
}
LEGEND_COLORS = ("FFFF0000", "theme1", "FF000000")

# PROJECTS sheet legend: gray = terminated, blue = quotation in process, green = for assignment of
# contractor, orange = on hold, yellow = for costing/quotation, fuchsia = quotation submitted.
PROJECT_SHEET_COLOR_STATUS = {
    "theme2": MARKETPLACE_LEAD_STATUS_TERMINATED,
    "FFBFBFBF": MARKETPLACE_LEAD_STATUS_TERMINATED,
    "FFD9D9D9": MARKETPLACE_LEAD_STATUS_TERMINATED,
    "FFA6A6A6": MARKETPLACE_LEAD_STATUS_TERMINATED,
    "FF808080": MARKETPLACE_LEAD_STATUS_TERMINATED,
    "FF00B0F0": MARKETPLACE_LEAD_STATUS_QUOTATION_IN_PROCESS,
    "FF0070C0": MARKETPLACE_LEAD_STATUS_QUOTATION_IN_PROCESS,
    "FF9BC2E6": MARKETPLACE_LEAD_STATUS_QUOTATION_IN_PROCESS,
    "FF92D050": MARKETPLACE_LEAD_STATUS_FOR_CONTRACTOR,
    "FF2DFF8C": MARKETPLACE_LEAD_STATUS_FOR_CONTRACTOR,
    "FF26DC18": MARKETPLACE_LEAD_STATUS_FOR_CONTRACTOR,
    "FF00B050": MARKETPLACE_LEAD_STATUS_FOR_CONTRACTOR,
    "FF00FF00": MARKETPLACE_LEAD_STATUS_FOR_CONTRACTOR,
    "theme5": MARKETPLACE_LEAD_STATUS_ON_HOLD,
    "FFFFC000": MARKETPLACE_LEAD_STATUS_ON_HOLD,
    "FFED7D31": MARKETPLACE_LEAD_STATUS_ON_HOLD,
    "FFF4B183": MARKETPLACE_LEAD_STATUS_ON_HOLD,
    "FFFFFF00": MARKETPLACE_LEAD_STATUS_FOR_QUOTATION,
    "FFFF66CC": MARKETPLACE_LEAD_STATUS_QUOTATION_SUBMITTED,
    "FFFF00FF": MARKETPLACE_LEAD_STATUS_QUOTATION_SUBMITTED,
}

# How each monitoring sheet is read. The PRODUCTS sheet types dates as MM.DD.YY text (Excel-converted
# cells come out day/month swapped); the PROJECTS sheet mostly stores real dates, so ambiguous cells are
# resolved against the sheet's other dates. Its color legend covers every status, so color wins there.
SHEET_PROFILES = {
    MARKETPLACE_CATEGORY_PRODUCTS: {
        "sheet_keyword": "PRODUCT",
        "colors": SHEET_COLOR_STATUS,
        "color_first": False,
        "date_mode": "swap",
        "status_text_fields": ("status_detail",),
    },
    TRANSACTION_TYPE_PROJECTS: {
        "sheet_keyword": "PROJECT",
        "colors": PROJECT_SHEET_COLOR_STATUS,
        "color_first": True,
        "date_mode": "auto",
        "status_text_fields": ("status_detail", "remarks", "action_needed"),
    },
}


def sheet_profile(transaction_type):
    return SHEET_PROFILES.get(transaction_type) or SHEET_PROFILES[MARKETPLACE_CATEGORY_PRODUCTS]

# Export fill per status so a downloaded sheet stays color coded.
STATUS_EXPORT_FILL = {
    "new": None,
    "for_research": "FFFFC000",
    "for_quotation": "FFFFFF00",
    "quotation_in_process": "FF00B0F0",
    "quotation_submitted": "FFFF00FF",
    "for_follow_up": "FF7030A0",
    "for_contractor": "FF26DC18",
    "on_hold": "FFF4B183",
    "awarded": "FF2E75B6",
    "mobilization": "FF8EA9DB",
    "ongoing": "FF70AD47",
    "turnover": "FFA9D08E",
    "order_confirmed": "FF2E75B6",
    "for_delivery": "FF8EA9DB",
    "delivered": "FFA9D08E",
    "completed": "FFFF0000",
    "terminated": "FF000000",
}

SHEET_HEADERS = (
    ("item_name", "TRANSACTION"),
    ("inquiry_no", "INQUIRYNO."),
    ("date_requested", "DATE REQUESTED"),
    ("estimated_implementation", "ESTIMATED DATE OF IMPLEMENTATION (EDI)"),
    ("aging", "AGING (DAYS) TODATE"),
    ("quantity", "QTY/ Volume"),
    ("specifications", "SPECIFICATIONS"),
    ("delivery_location", "Delivered to/Location"),
    ("referrer_name", "PROF"),
    ("referrer_phone", "PROF'S CONTACT NUMBER"),
    ("referrer_email", "PROF EMAIL"),
    ("client_company", "CLIENT / COMPANY"),
    ("guest_name", "CLIENT'S REPRESENTATIVE"),
    ("guest_phone", "CLIENT'S CONTACT NUMBER"),
    ("guest_email", "CLIENT'S EMAIL"),
    ("status_detail", "STATUS"),
    ("action_needed", "ACTION NEEDED"),
    ("project_coordinator", "PROJECT COORDINATOR"),
    ("assigned_contractor", "Assigned Contractor"),
    ("date_completed", "DATE COMPLETED"),
    ("remarks", "REMARKS"),
    ("reference_number", "REFERENCE NO."),
    ("workflow_status", "WORKFLOW STATUS"),
    ("assigned_staff", "ASSIGNED STAFF"),
)

PROJECT_SHEET_HEADER_OVERRIDES = {
    "inquiry_no": "No.",
    "specifications": "PROJECT DETAILS AND SPECIFICATIONS",
    "delivery_location": "LOCATION",
}


def sheet_headers(transaction_type=None):
    if transaction_type != TRANSACTION_TYPE_PROJECTS:
        return SHEET_HEADERS
    return tuple(
        (field, PROJECT_SHEET_HEADER_OVERRIDES.get(field, header))
        for field, header in SHEET_HEADERS
        if field != "quantity"
    )


MULTILINE_FIELDS = ("quantity", "specifications")
PHONE_FIELDS = ("referrer_phone", "guest_phone")


def _norm_header(value):
    return re.sub(r"[^A-Z0-9]", "", str(value or "").upper())


def _header_field(header):
    h = _norm_header(header)
    if not h:
        return None
    if h.startswith("INQUIRYNO") or h == "NO":
        return "inquiry_no"
    if h == "TRANSACTION":
        return "item_name"
    if h.startswith("DATEREQUESTED"):
        return "date_requested"
    if h.startswith("ESTIMATEDDATE") or h == "EDI":
        return "estimated_implementation"
    if h.startswith("AGING"):
        return "aging"
    if h.startswith("QTY") or h.startswith("QUANTITY"):
        return "quantity"
    if h.startswith("SPECIFICATION") or h.startswith("PROJECTDETAILS"):
        return "specifications"
    if h.startswith("DELIVEREDTO") or h == "LOCATION":
        return "delivery_location"
    if h == "PROF":
        return "referrer_name"
    if h.startswith("PROF") and "CONTACT" in h:
        return "referrer_phone"
    if h.startswith("PROF") and "EMAIL" in h:
        return "referrer_email"
    if h.startswith("CLIENTCOMPANY") or h == "COMPANY":
        return "client_company"
    if h.startswith("CLIENT") and "REPRESENTATIVE" in h:
        return "guest_name"
    if h.startswith("CLIENT") and "CONTACT" in h:
        return "guest_phone"
    if "CLIENT" in h and "EMAIL" in h:
        return "guest_email"
    if h == "STATUS":
        return "status_detail"
    if h == "WORKFLOWSTATUS":
        return "workflow_status"
    # The PROJECTS sheet splits "ACTION / NEEDED" across two header rows.
    if h.startswith("ACTIONNEEDED") or h == "ACTION":
        return "action_needed"
    if h.startswith("PROJECTCOORDINATOR"):
        return "project_coordinator"
    if h.startswith("ASSIGNEDCONTRACTOR"):
        return "assigned_contractor"
    if h.startswith("DATECOMPLETED"):
        return "date_completed"
    if h.startswith("REMARK"):
        return "remarks"
    if h.startswith("REFERENCENO"):
        return "reference_number"
    return None


def _cell_text(value, phone=False):
    if value is None:
        return None
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, (int, float)):
        if float(value).is_integer():
            text = str(int(value))
            # Excel drops the leading zero of PH mobile numbers (09xx...).
            if phone and len(text) == 10 and text.startswith("9"):
                text = "0" + text
            return text
        return str(value)
    if isinstance(value, datetime):
        return value.strftime("%m.%d.%y")
    if isinstance(value, date):
        return value.strftime("%m.%d.%y")
    text = str(value).strip()
    if phone and re.match(r"9\d{9}(?!\d)", text):
        text = "0" + text
    return text or None


def _swapped(value):
    try:
        return date(value.year, value.day, value.month)
    except ValueError:
        return None


def _cell_date(value, mode="swap", reference=None):
    """
    Parse a sheet date. mode "swap": Excel-converted cells are day/month swapped (PRODUCTS sheet).
    mode "auto": ambiguous cells (day <= 12) take whichever reading is closer to `reference`.
    """
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        as_is = value.date()
        swapped = _swapped(value) if value.day <= 12 else None
        if swapped is None:
            return as_is
        if mode == "swap":
            return swapped
        if reference and abs((swapped - reference).days) < abs((as_is - reference).days):
            return swapped
        return as_is
    if isinstance(value, date):
        return value
    text = str(value).strip()
    match = re.fullmatch(r"(\d{1,2})[./-](\d{1,2})(?:[./-](\d{2,4}))?", text)
    if match:
        month, day = int(match.group(1)), int(match.group(2))
        if match.group(3):
            year = int(match.group(3))
            if year < 100:
                year += 2000
        else:
            year = (reference or manila_today()).year
        try:
            return date(year, month, day)
        except ValueError:
            return None
    try:
        return datetime.strptime(text[:10], "%Y-%m-%d").date()
    except ValueError:
        return None


def _reference_date(sheet, header_row, column):
    """Median of the unambiguous request dates on the sheet (used to resolve day/month mix-ups)."""
    known = []
    for row_idx in range(header_row + 1, sheet.max_row + 1):
        value = sheet.cell(row_idx, column).value
        if isinstance(value, datetime):
            if value.day > 12:
                known.append(value.date())
        elif isinstance(value, str) and re.fullmatch(r"\d{1,2}[./-]\d{1,2}[./-]\d{2,4}", value.strip()):
            parsed = _cell_date(value)
            if parsed:
                known.append(parsed)
    if not known:
        return None
    known.sort()
    return known[len(known) // 2]


def _fill_key(cell):
    fill = cell.fill
    if fill is None or fill.fill_type != "solid":
        return None
    color = fill.fgColor
    if color is None:
        return None
    if color.type == "rgb" and isinstance(color.rgb, str):
        return color.rgb.upper()
    if color.type == "theme":
        return f"theme{color.theme}"
    return None


def status_from_sheet(status_text, color_key, profile=None):
    """Map the sheet's free-text STATUS plus row color to a workflow status."""
    profile = profile or SHEET_PROFILES[MARKETPLACE_CATEGORY_PRODUCTS]
    colors = profile["colors"]
    text = (status_text or "").upper()
    if any(word in text for word in ("TERMINATED", "SCAM", "CANCEL", "DID NOT WIN", "DEAD")):
        return MARKETPLACE_LEAD_STATUS_TERMINATED
    if profile["color_first"] and colors.get(color_key):
        return colors[color_key]
    if not profile["color_first"] and color_key in LEGEND_COLORS:
        return SHEET_COLOR_STATUS[color_key]
    if "COMPLETED" in text:
        return MARKETPLACE_LEAD_STATUS_COMPLETED
    if "RESEARCH" in text:
        return MARKETPLACE_LEAD_STATUS_FOR_RESEARCH
    if any(word in text for word in ("QUOTATION SUBMITTED", "SUBMITTED QUOTATION", "WITH QUOTATION")):
        return MARKETPLACE_LEAD_STATUS_QUOTATION_SUBMITTED
    if re.search(r"QUOTATION\s+(ON[\s-]?GOING|IN[\s-]?PROCESS)", text):
        return MARKETPLACE_LEAD_STATUS_QUOTATION_IN_PROCESS
    if re.search(r"ASSIGN\w*\s+OF\s+CONTRACTOR", text):
        return MARKETPLACE_LEAD_STATUS_FOR_CONTRACTOR
    if "QUOTATION" in text or "COSTING" in text:
        return MARKETPLACE_LEAD_STATUS_FOR_QUOTATION
    if "ON HOLD" in text or "WAITING" in text:
        return MARKETPLACE_LEAD_STATUS_ON_HOLD
    if re.search(r"FOLLOW[\s-]?UP|\bCALL\b", text):
        return MARKETPLACE_LEAD_STATUS_FOR_FOLLOW_UP
    if colors.get(color_key):
        return colors[color_key]
    return MARKETPLACE_LEAD_STATUS_NEW


def _status_from_label(label):
    text = (label or "").strip().lower()
    if not text:
        return None
    for key, value in MARKETPLACE_LEAD_STATUS_LABELS.items():
        if text in (key, value.lower()):
            return key
    return None


def _member_name_index():
    index = {}
    for member in Member.query.all():
        names = {member.full_name}
        names.add(" ".join(p for p in (member.first_name, member.last_name) if p))
        for name in names:
            key = re.sub(r"[^a-z]", "", (name or "").lower())
            if key:
                index.setdefault(key, set()).add(member.member_id)
    return index


def _match_member(index, name):
    key = re.sub(r"[^a-z]", "", (name or "").lower())
    ids = index.get(key) if key else None
    if ids and len(ids) == 1:
        return next(iter(ids))
    return None


def _pick_sheet(workbook, sheet_name=None, keyword="PRODUCT"):
    if sheet_name:
        if sheet_name not in workbook.sheetnames:
            raise ValueError(f"Sheet '{sheet_name}' was not found in the workbook.")
        return workbook[sheet_name]
    for name in workbook.sheetnames:
        if keyword in name.upper():
            return workbook[name]
    return workbook.active


def _locate_header(sheet):
    for row_idx in range(1, min(sheet.max_row, 20) + 1):
        columns = {}
        for col_idx in range(1, min(sheet.max_column, 60) + 1):
            field = _header_field(sheet.cell(row_idx, col_idx).value)
            if field and field not in columns:
                columns[field] = col_idx
        if "inquiry_no" in columns and "item_name" in columns:
            return row_idx, columns
    raise ValueError(
        "Could not find the header row. The sheet needs TRANSACTION and INQUIRY NO. (or No.) columns "
        "(same layout as the TBGP monitoring sheet)."
    )


def _parse_sheet_records(sheet, profile=None):
    profile = profile or SHEET_PROFILES[MARKETPLACE_CATEGORY_PRODUCTS]
    header_row, columns = _locate_header(sheet)
    date_mode = profile["date_mode"]
    reference = (
        _reference_date(sheet, header_row, columns["date_requested"])
        if date_mode == "auto" and "date_requested" in columns
        else None
    )

    def parse_date(value):
        return _cell_date(value, mode=date_mode, reference=reference)

    records = []
    current = None
    for row_idx in range(header_row + 1, sheet.max_row + 1):
        values = {field: sheet.cell(row_idx, col).value for field, col in columns.items()}
        if not any(v not in (None, "") for v in values.values()):
            continue
        starts_record = values.get("inquiry_no") not in (None, "") or values.get("item_name") not in (None, "")
        if starts_record:
            current = {
                "row": row_idx,
                "color": _fill_key(sheet.cell(row_idx, columns["item_name"])),
                "raw": {},
                "dates": {},
            }
            records.append(current)
        elif current is None:
            continue
        for field, value in values.items():
            if field in ("date_requested", "date_completed"):
                if starts_record or field not in current["dates"]:
                    parsed = parse_date(value)
                    if parsed:
                        current["dates"][field] = parsed
                if starts_record and field == "date_requested" and value not in (None, "") and not parse_date(value):
                    current["raw"].setdefault("remarks", []).append(f"Date requested: {value}")
                continue
            if field == "aging":
                continue
            text = _cell_text(value, phone=field in PHONE_FIELDS)
            if not text:
                continue
            bucket = current["raw"].setdefault(field, [])
            if field in MULTILINE_FIELDS or text not in bucket:
                bucket.append(text)
    return records


def _finalize_record(record):
    data = {}
    for field, parts in record["raw"].items():
        joiner = "\n" if field in MULTILINE_FIELDS else " / "
        data[field] = joiner.join(parts)
    data.update(record["dates"])
    return data


def _apply_import_fields(lead, data):
    for field, (_label, max_len) in TRANSACTION_TEXT_FIELDS.items():
        if field == "message":
            continue
        if field in data:
            setattr(lead, field, _clean(data[field], max_len))
    lead.date_requested = data.get("date_requested") or lead.date_requested
    lead.date_completed = data.get("date_completed") or lead.date_completed


def _renumber_legacy(lead, floor, user_id):
    """Move a pre-reference-number CRM inquiry out of the way of the sheet's numbering."""
    old_ref = lead.reference_number
    lead.inquiry_no = max(next_inquiry_no(lead.transaction_type), floor)
    lead.reference_number = format_reference_number(lead.transaction_type, lead.inquiry_no)
    db.session.flush()
    _append_history(
        lead,
        event_type="update",
        note=f"Reference renumbered from {old_ref} to {lead.reference_number} for the Excel import",
        user_id=user_id,
    )


def import_transactions_from_xlsx(
    source,
    *,
    transaction_type=MARKETPLACE_CATEGORY_PRODUCTS,
    update_existing=False,
    sheet_name=None,
    user_id=None,
):
    """
    Import the TBGP color-coded monitoring sheet into transactions.

    Records are keyed by (transaction type, inquiry no.). Existing imported rows are skipped
    unless update_existing is set; numbers held by website/manual transactions are reported.
    """
    if transaction_type not in TRANSACTION_TYPE_LABELS:
        raise ValueError("Invalid transaction type.")
    try:
        workbook = load_workbook(source, data_only=True)
    except Exception as exc:  # noqa: BLE001 - openpyxl raises many types for bad files
        raise ValueError(f"Could not read the Excel file: {exc}") from exc
    profile = sheet_profile(transaction_type)
    sheet = _pick_sheet(workbook, sheet_name, profile["sheet_keyword"])
    records = _parse_sheet_records(sheet, profile)
    member_index = _member_name_index()
    coordinators = duty_holders(DUTY_PROJECT_COORDINATOR) if transaction_type == TRANSACTION_TYPE_PROJECTS else []
    sheet_numbers = [
        int(text) for text in ((r["raw"].get("inquiry_no") or [""])[0] for r in records) if text.isdigit()
    ]
    renumber_floor = (max(sheet_numbers) if sheet_numbers else 0) + 1

    summary = {"sheet": sheet.title, "created": 0, "updated": 0, "skipped": 0, "errors": []}
    for record in records:
        data = _finalize_record(record)
        inquiry_raw = (data.pop("inquiry_no", "") or "").strip()
        if not inquiry_raw.isdigit():
            summary["errors"].append(f"Row {record['row']}: missing or invalid INQUIRY NO. ({inquiry_raw or 'blank'})")
            continue
        inquiry_no = int(inquiry_raw)
        workflow = _status_from_label(data.pop("workflow_status", None))
        status_text = " | ".join(data[f] for f in profile["status_text_fields"] if data.get(f))
        status = workflow or status_from_sheet(status_text, record["color"], profile)
        sheet_reference = (data.pop("reference_number", "") or "").strip().upper()
        data.pop("assigned_staff", None)

        existing = MarketplaceLead.query.filter_by(
            transaction_type=transaction_type, inquiry_no=inquiry_no
        ).first()
        # A row exported from the portal carries its reference no. and is the same record.
        same_record = existing is not None and sheet_reference == (existing.reference_number or "").upper()
        if existing is not None and not same_record and existing.source == TRANSACTION_SOURCE_LEGACY:
            _renumber_legacy(existing, renumber_floor, user_id)
            existing = None
        if existing is not None and not same_record and existing.source != TRANSACTION_SOURCE_IMPORT:
            summary["errors"].append(
                f"Row {record['row']}: inquiry no. {inquiry_no} is already used by "
                f"{existing.reference_number} ({transaction_source_label(existing.source)}); skipped."
            )
            continue
        if existing is not None and not update_existing:
            summary["skipped"] += 1
            continue

        member_id = _match_member(member_index, data.get("referrer_name"))
        now = manila_now()
        if existing is None:
            requested = data.get("date_requested")
            lead = MarketplaceLead(
                transaction_type=transaction_type,
                interest_category=transaction_type,
                inquiry_no=inquiry_no,
                reference_number=format_reference_number(transaction_type, inquiry_no),
                source=TRANSACTION_SOURCE_IMPORT,
                status=status,
                created_at=datetime.combine(requested, time()) if requested else now,
                updated_at=now,
                status_changed_at=now,
            )
            _apply_import_fields(lead, data)
            lead.attributed_member_id = member_id
            coordinator = match_coordinator(lead.project_coordinator, coordinators) if coordinators else None
            lead.assigned_user_id = coordinator.user_id if coordinator else None
            if not lead.item_name:
                lead.item_name = "(unnamed)"
            db.session.add(lead)
            db.session.flush()
            _append_history(
                lead,
                event_type="import",
                note=f"Imported from Excel ({sheet.title} row {record['row']})",
                user_id=user_id,
            )
            summary["created"] += 1
        else:
            previous_status = existing.status
            _apply_import_fields(existing, data)
            if previous_status != status:
                existing.status_changed_at = now
            existing.status = status
            if member_id and not existing.attributed_member_id:
                existing.attributed_member_id = member_id
            if coordinators and not existing.assigned_user_id:
                coordinator = match_coordinator(existing.project_coordinator, coordinators)
                existing.assigned_user_id = coordinator.user_id if coordinator else None
            existing.updated_at = now
            note = f"Updated from Excel ({sheet.title} row {record['row']})"
            if previous_status != status:
                note += (
                    f"; Status: {transaction_status_label(previous_status)} -> "
                    f"{transaction_status_label(status)}"
                )
            _append_history(existing, event_type="import", note=note, user_id=user_id)
            summary["updated"] += 1

    db.session.commit()
    return summary


# --------------------------------------------------------------------------- Excel export


def export_transactions_xlsx(leads, transaction_type=None):
    """Monitoring-sheet layout (re-importable), color coded by workflow status."""
    headers = sheet_headers(transaction_type)
    workbook = Workbook()
    sheet = workbook.active
    if transaction_type == TRANSACTION_TYPE_PROJECTS:
        sheet.title = "PROJECTS"
        sheet.cell(1, 1, "COLOR CODING:")
        sheet.cell(1, 2, ", ".join(
            f"{transaction_status_label(status)}" for status in PROJECT_LEAD_STATUSES
        ) + " (row color follows WORKFLOW STATUS)")
    else:
        sheet.title = "PRODUCTS"
        sheet.cell(1, 1, manila_now().strftime("%m.%d.%y"))
        sheet.cell(1, 3, "RED-COMPLETED")
        sheet.cell(1, 5, "BLACK -DEAD")
    header_fill = PatternFill("solid", fgColor="FFE2EFDA")
    for col, (_field, header) in enumerate(headers, start=1):
        cell = sheet.cell(2, col, header)
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        sheet.column_dimensions[get_column_letter(col)].width = 22

    for row_idx, lead in enumerate(leads, start=3):
        values = {
            "item_name": lead.display_item,
            "inquiry_no": lead.inquiry_no,
            "date_requested": lead.date_requested.strftime("%m.%d.%y") if lead.date_requested else None,
            "aging": lead.aging_days,
            "date_completed": lead.date_completed.strftime("%m.%d.%y") if lead.date_completed else None,
            "referrer_name": lead.referred_by_name or None,
            "workflow_status": transaction_status_label(lead.status),
            "assigned_staff": (
                (lead.assigned_user.full_name or lead.assigned_user.username) if lead.assigned_user else None
            ),
        }
        fill_rgb = STATUS_EXPORT_FILL.get(lead.status or "new")
        fill = PatternFill("solid", fgColor=fill_rgb) if fill_rgb else None
        light_text = fill_rgb in ("FFFF0000", "FF000000", "FF7030A0", "FF2E75B6")
        for col, (field, _header) in enumerate(headers, start=1):
            value = values[field] if field in values else getattr(lead, field, None)
            cell = sheet.cell(row_idx, col, value)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if fill is not None:
                cell.fill = fill
                if light_text:
                    cell.font = Font(color="FFFFFFFF")
    sheet.freeze_panes = "C3"

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output
